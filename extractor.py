# extractors.py
import fitz
import tempfile
import os
from pdf2image import convert_from_path
from PIL import Image
from pytesseract import pytesseract, Output
from docx import Document
from spire.xls import *
from spire.xls.common import *

def extract_text_from_pdf(pdf_file):
    try:
        doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
        text = "".join([page.get_text() for page in doc])
        if not text.strip():
            print("Searchable text not found in PDF. Using OCR to extract text.")
            pdf_file.seek(0)
            text = extract_text_with_ocr(pdf_file)
        return text
    except Exception as e:
        return f"Error extracting text from PDF: {e}"

def extract_text_with_ocr(pdf_file):
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp_pdf:
            temp_pdf.write(pdf_file.read())
            temp_pdf_path = temp_pdf.name

        images = convert_from_path(temp_pdf_path, dpi=300, fmt="jpeg")
        text = "".join(
            pytesseract.image_to_string(image, lang="eng", output_type=Output.STRING)
            for image in images
        )
        return text
    except Exception as e:
        return f"Error performing OCR on PDF: {e}"

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    return "\n".join([paragraph.text for paragraph in doc.paragraphs])

def extract_data_from_excel(excel_file):
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_excel:
            temp_excel.write(excel_file.read())
            temp_excel_path = temp_excel.name

        workbook = Workbook()
        workbook.LoadFromFile(temp_excel_path)

        sheet = workbook.Worksheets[0]
        temp_txt_path = tempfile.NamedTemporaryFile(delete=False, suffix=".txt").name
        sheet.SaveToFile(temp_txt_path, '\t', Encoding.get_UTF8())

        with open(temp_txt_path, 'r', encoding='utf-8') as txt_file:
            extracted_text = txt_file.read()

        os.remove(temp_txt_path)
        workbook.Dispose()

        return extracted_text
    except Exception as e:
        return f"Error extracting data from Excel: {e}"

def extract_text_from_image(image_file):
    """
    Processes an image file using both BLIP and Tesseract OCR.
    BLIP is used to generate a contextual caption, and OCR extracts any text in the image.
    """
    try:
        # Open the image
        image = Image.open(image_file).convert("RGB")
    except Exception as e:
        return f"Error opening image: {e}"

    # Initialize variables to store results
    ocr_text = ""
    blip_caption = ""

    # Extract text using OCR (Tesseract)
    try:
        ocr_text = pytesseract.image_to_string(image, lang="eng", output_type=Output.STRING)
    except Exception as e:
        ocr_text = f"Error extracting OCR text: {e}"

    # Extract context using BLIP
    try:
        from transformers import BlipProcessor, BlipForConditionalGeneration
        processor = BlipProcessor.from_pretrained("Salesforce/blip-image-captioning-base")
        model = BlipForConditionalGeneration.from_pretrained("Salesforce/blip-image-captioning-base")
        inputs = processor(image, return_tensors="pt")
        out = model.generate(**inputs)
        blip_caption = processor.decode(out[0], skip_special_tokens=True)
    except Exception as e:
        blip_caption = f"Error generating BLIP caption: {e}"

    # Combine results
    combined_text = f"BLIP Caption:\n{blip_caption}\n\nOCR Extracted Text:\n{ocr_text}"
    return combined_text
